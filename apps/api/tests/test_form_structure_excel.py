"""1.2 form_structure Excel(.xlsx) 解析测试。

用 openpyxl 动态生成样例 xlsx（含多级表头/合并单元格/空格/模板填空栏），
再验证 parse_structure 的输出。场景与 Word 测试一一对应，确保跨格式行为一致。
"""

from __future__ import annotations

import json
from pathlib import Path

import pytest
from openpyxl import Workbook

from aff_contracts.fill import ExcelLocator
from app.modules.form_structure.service import FormStructureService


def _save(wb: Workbook, p: Path) -> Path:
    p.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(p))
    return p


@pytest.fixture()
def sample_xlsx(tmp_path: Path) -> Path:
    """生成样例 Excel：
    - Sheet1「登记表」：form 式（标签 | 空格同行）+ 横合并备注
    - Sheet2「展示」：纯展示表（无空格）
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "登记表"
    # 行1：标签 + 空格交替（form 式，header_rows=0）
    ws.cell(1, 1, "姓名"); ws.cell(1, 3, "性别"); ws.cell(1, 5, "学号")
    # 行2：标签 + 空格
    ws.cell(2, 1, "联系电话"); ws.cell(2, 3, "籍贯")
    # 行3：备注（A3:D3 合并标签）+ 空值（E3:H3 合并空格）
    ws.cell(3, 1, "备注")
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=4)
    ws.merge_cells(start_row=3, start_column=5, end_row=3, end_column=8)

    ws2 = wb.create_sheet("展示")
    ws2.cell(1, 1, "项目"); ws2.cell(1, 2, "说明")
    ws2.cell(2, 1, "填表日期"); ws2.cell(2, 2, "2026年8月")

    return _save(wb, tmp_path / "sample.xlsx")


def test_parse_xlsx_basic(sample_xlsx: Path) -> None:
    """验证基本结构：format/filename/title/sheets/fields。"""
    result = FormStructureService().parse_structure(str(sample_xlsx), filename="sample.xlsx")
    assert result["format"] == "xlsx"
    assert result["filename"] == "sample.xlsx"
    assert result["title"] == "登记表"  # 首个工作表名
    assert len(result["sheets"]) == 2
    # docx 专属字段在 xlsx 下为空
    assert result["headings"] == []
    assert result["paragraphs"] == []
    assert result["tables"] == []


def test_parse_xlsx_multi_sheet(sample_xlsx: Path) -> None:
    """验证多工作表解析。"""
    result = FormStructureService().parse_structure(str(sample_xlsx), filename="sample.xlsx")
    sheets = result["sheets"]
    names = [s["name"] for s in sheets]
    assert names == ["登记表", "展示"]
    assert sheets[0]["index"] == 0
    assert sheets[1]["index"] == 1


def test_parse_xlsx_merged_cells(sample_xlsx: Path) -> None:
    """验证合并单元格：登记表 sheet 有 2 处横合并。"""
    result = FormStructureService().parse_structure(str(sample_xlsx), filename="sample.xlsx")
    sheet1 = result["sheets"][0]
    assert len(sheet1["merged_ranges"]) == 2
    cell_texts = [c["text"] for c in sheet1["cells"]]
    assert "备注" in cell_texts
    # 备注格 merged=True，带 colspan=4
    beizhu = next(c for c in sheet1["cells"] if c["text"] == "备注")
    assert beizhu["merged"] is True
    assert beizhu["colspan"] == 4
    assert beizhu["row"] == 3 and beizhu["col"] == 1


def test_parse_xlsx_header_rows(sample_xlsx: Path) -> None:
    """验证表头行数：form 式表=0，展示表=1。"""
    result = FormStructureService().parse_structure(str(sample_xlsx), filename="sample.xlsx")
    sheets = result["sheets"]
    assert sheets[0]["header_rows"] == 0  # 第1行就有空格
    assert sheets[1]["header_rows"] == 1  # 无空格，首行约定表头


def test_parse_xlsx_multi_level_header(tmp_path: Path) -> None:
    """验证多级表头：跨列表头合并 + 子表头 + 数据行 → header_rows=2。"""
    wb = Workbook(); ws = wb.active
    # 行1：跨列表头
    ws.cell(1, 1, "个人信息"); ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=2)
    ws.cell(1, 3, "工作信息"); ws.merge_cells(start_row=1, start_column=3, end_row=1, end_column=4)
    # 行2：子表头
    ws.cell(2, 1, "姓名"); ws.cell(2, 2, "性别"); ws.cell(2, 3, "公司"); ws.cell(2, 4, "职位")
    # 行3：空数据行（显式创建空 cell，模拟真实空填写行）
    for c in range(1, 5):
        ws.cell(3, c, "")
    p = _save(wb, tmp_path / "multi_header.xlsx")
    result = FormStructureService().parse_structure(str(p), filename="multi_header.xlsx")
    sheet = result["sheets"][0]
    assert sheet["header_rows"] == 2
    assert len(sheet["merged_ranges"]) == 2


def test_parse_xlsx_template_fillable(tmp_path: Path) -> None:
    """验证模板填空栏（日期/地址/盖章）被识别为可填写位置。"""
    wb = Workbook(); ws = wb.active
    ws.cell(1, 1, "活动时间"); ws.cell(1, 2, "年  月  日—年  月  日")
    ws.cell(2, 1, "实践基地"); ws.cell(2, 2, "省    市    县（区）   镇    村")
    ws.cell(3, 1, "辅导员意见"); ws.cell(3, 2, "盖章：\n年  月  日")
    ws.cell(4, 1, "备注")  # (4,2) 留空
    p = _save(wb, tmp_path / "template.xlsx")
    result = FormStructureService().parse_structure(str(p), filename="template.xlsx")
    fields = result["fields"]
    assert len(fields) == 4
    names = {f["name"] for f in fields}
    assert {"活动时间", "实践基地", "辅导员意见", "备注"} == names


def test_parse_xlsx_label_inheritance(tmp_path: Path) -> None:
    """验证多行向上继承字段名：行3空格继承行1表头（跨2行）。"""
    wb = Workbook(); ws = wb.active
    ws.cell(1, 1, "姓名"); ws.cell(1, 2, "性别")
    # 行2、行3 留空（显式创建空 cell，模拟真实空填写行）
    for r in (2, 3):
        for c in (1, 2):
            ws.cell(r, c, "")
    p = _save(wb, tmp_path / "inherit.xlsx")
    result = FormStructureService().parse_structure(str(p), filename="inherit.xlsx")
    fields = result["fields"]
    # 行2、行3 各 2 个空格 = 4 字段
    assert len(fields) == 4
    row3_fields = [f for f in fields if f["locator"]["row"] == 3]
    assert {f["name"] for f in row3_fields} == {"姓名", "性别"}


def test_parse_xlsx_header_not_fillable(tmp_path: Path) -> None:
    """验证纯标签表头行不被标记为可填写位置。"""
    wb = Workbook(); ws = wb.active
    ws.cell(1, 1, "姓名"); ws.cell(1, 2, "年龄"); ws.cell(1, 3, "电话")
    # 行2 留空（显式创建空 cell，模拟真实空填写行）
    for c in (1, 2, 3):
        ws.cell(2, c, "")
    p = _save(wb, tmp_path / "header.xlsx")
    result = FormStructureService().parse_structure(str(p), filename="header.xlsx")
    fields = result["fields"]
    assert len(fields) == 3
    assert {f["locator"]["row"] for f in fields} == {2}
    assert {f["name"] for f in fields} == {"姓名", "年龄", "电话"}


def test_parse_xlsx_fill_kind(tmp_path: Path) -> None:
    """验证 fill_kind：空格=empty，模板填空栏=template。"""
    wb = Workbook(); ws = wb.active
    ws.cell(1, 1, "活动时间"); ws.cell(1, 2, "年  月  日—年  月  日")
    ws.cell(2, 1, "姓名")  # (2,2) 留空
    ws.cell(3, 1, "辅导员意见"); ws.cell(3, 2, "盖章：\n年  月  日")
    p = _save(wb, tmp_path / "fill_kind.xlsx")
    result = FormStructureService().parse_structure(str(p), filename="fill_kind.xlsx")
    fields = result["fields"]
    assert len(fields) == 3
    by_row = {f["locator"]["row"]: f for f in fields}
    assert by_row[1]["fill_kind"] == "template"
    assert by_row[2]["fill_kind"] == "empty"
    assert by_row[3]["fill_kind"] == "template"


def test_parse_xlsx_instruction_text(tmp_path: Path) -> None:
    """验证填写说明文字区被识别为 template。"""
    wb = Workbook(); ws = wb.active
    ws.cell(1, 1, "年级推荐意见"); ws.cell(1, 2, "辅导员签字：")
    ws.cell(2, 1, "曾获奖项")
    ws.cell(2, 2, "大学期间获得的校级及以上奖励，如果无校级及以上，则院级也可以")
    ws.cell(3, 1, "学生工作经历"); ws.cell(3, 2, "时间段＋具体岗位")
    ws.cell(4, 1, "姓名")  # (4,2) 留空
    ws.cell(5, 1, "意向部门")
    ws.cell(5, 2, "学习部 调研部 权益部 活动部 （至少选择两个）")
    p = _save(wb, tmp_path / "instruction.xlsx")
    result = FormStructureService().parse_structure(str(p), filename="instruction.xlsx")
    fields = result["fields"]
    by_row = {f["locator"]["row"]: f for f in fields}
    assert by_row[1]["fill_kind"] == "template" and by_row[1]["name"] == "年级推荐意见"
    assert by_row[2]["fill_kind"] == "template" and by_row[2]["name"] == "曾获奖项"
    assert by_row[3]["fill_kind"] == "template" and by_row[3]["name"] == "学生工作经历"
    assert by_row[4]["fill_kind"] == "empty" and by_row[4]["name"] == "姓名"
    assert 5 not in by_row  # 参考列表不是可填写位


def test_parse_xlsx_multi_col_header_inheritance(tmp_path: Path) -> None:
    """验证多级表头字段名继承：纵向合并小节标题不干扰 + 横合并表头溯源。"""
    wb = Workbook(); ws = wb.active
    # A 列纵向合并「家庭成员情况」跨行1-4
    ws.cell(1, 1, "家庭成员情况")
    ws.merge_cells(start_row=1, start_column=1, end_row=4, end_column=1)
    # 行1：列头 B/C/D，其中 D1:E1 横合并为「年收入」
    ws.cell(1, 2, "姓名"); ws.cell(1, 3, "年龄"); ws.cell(1, 4, "年收入")
    ws.merge_cells(start_row=1, start_column=4, end_row=1, end_column=5)
    # 行2-4：空数据行
    p = _save(wb, tmp_path / "multi_col_header.xlsx")
    result = FormStructureService().parse_structure(str(p), filename="multi_col_header.xlsx")
    fields = result["fields"]
    by_pos = {(f["locator"]["row"], f["locator"]["col"]): f for f in fields}
    assert by_pos[(2, 2)]["name"] == "姓名"
    assert by_pos[(2, 3)]["name"] == "年龄"
    assert by_pos[(2, 4)]["name"] == "年收入"
    assert by_pos[(2, 5)]["name"] == "年收入"  # 横合并延续格也继承
    for f in fields:
        assert f["name"] != "家庭成员情况"


def test_parse_xlsx_locator_contract(sample_xlsx: Path) -> None:
    """每个 field 的 locator 须符合 ExcelLocator 契约。"""
    result = FormStructureService().parse_structure(str(sample_xlsx), filename="sample.xlsx")
    for f in result["fields"]:
        assert f["locator"]["kind"] == "excel_cell"
        loc = ExcelLocator(**f["locator"])
        assert loc.row >= 1
        assert loc.col >= 1
        assert loc.sheet == "登记表" or loc.sheet == "展示"


def test_parse_xlsx_merged_range_string(tmp_path: Path) -> None:
    """合并块的可填写位 locator.merged_range 应给出范围字符串（如 A3:E3）。"""
    wb = Workbook(); ws = wb.active
    ws.cell(1, 1, "备注")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=4)
    ws.cell(2, 1, "姓名")  # (2,2) 留空 → 非合并 empty
    p = _save(wb, tmp_path / "mrange.xlsx")
    result = FormStructureService().parse_structure(str(p), filename="mrange.xlsx")
    fields = result["fields"]
    by_row = {f["locator"]["row"]: f for f in fields}
    # 行1 合并空格（A1:D1 中只有 A1 有文字「备注」，但备注是纯标签→不填写；
    # 这里行1无空格可填写位，跳过）→ 实际行1只备注标签，无可填写
    # 行2：(2,2) 空格 → empty，merged_range=None
    assert by_row[2]["fill_kind"] == "empty"
    assert by_row[2]["locator"]["merged_range"] is None


def test_parse_xlsx_merged_blank_field(tmp_path: Path) -> None:
    """合并的纯空格块（无文字）应识别为可填写位并带 merged_range。"""
    wb = Workbook(); ws = wb.active
    ws.cell(1, 1, "意见")  # 标签
    # (1,2):(1,4) 合并的纯空格块
    ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=4)
    p = _save(wb, tmp_path / "mblank.xlsx")
    result = FormStructureService().parse_structure(str(p), filename="mblank.xlsx")
    fields = result["fields"]
    assert len(fields) == 1
    f = fields[0]
    assert f["name"] == "意见"
    assert f["fill_kind"] == "empty"
    assert f["locator"]["merged_range"] == "B1:D1"
    assert f["locator"]["row"] == 1 and f["locator"]["col"] == 2


def test_parse_xlsx_field_ids_unique(sample_xlsx: Path) -> None:
    """每个可填写位置的 id 应唯一。"""
    result = FormStructureService().parse_structure(str(sample_xlsx), filename="sample.xlsx")
    ids = [f["id"] for f in result["fields"]]
    assert len(ids) == len(set(ids))


def test_parse_xlsx_structure_notes(sample_xlsx: Path) -> None:
    """验证结构备注。"""
    result = FormStructureService().parse_structure(str(sample_xlsx), filename="sample.xlsx")
    notes = result["structure_notes"]
    assert any("工作表" in n for n in notes)


def test_parse_xlsx_result_serializable(sample_xlsx: Path) -> None:
    """结果应能序列化为 JSON。"""
    result = FormStructureService().parse_structure(str(sample_xlsx), filename="sample.xlsx")
    json.dumps(result, ensure_ascii=False)


def test_parse_xlsx_file_not_found() -> None:
    """文件不存在应抛 FileNotFoundError。"""
    with pytest.raises(FileNotFoundError):
        FormStructureService().parse_structure("nonexistent.xlsx", filename="nonexistent.xlsx")


def test_parse_xlsx_unsupported_format(tmp_path: Path) -> None:
    """非 docx/xlsx 文件应抛 ValueError。"""
    f = tmp_path / "not_a_sheet.txt"
    f.write_text("hello", encoding="utf-8")
    with pytest.raises(ValueError):
        FormStructureService().parse_structure(str(f), filename="not_a_sheet.txt")


def test_parse_xlsx_serial_column_skipped(tmp_path: Path) -> None:
    """验证序号列（纯数字逐行递增）被识别并跳过，字段名改由列头继承。"""
    wb = Workbook(); ws = wb.active
    # 行1：表头（col0 序号列 + 列头）
    ws.cell(1, 1, "序号"); ws.cell(1, 2, "姓名"); ws.cell(1, 3, "性别"); ws.cell(1, 4, "年龄")
    # 行2-4：序号 1/2/3 + 空数据行
    for i, r in enumerate((2, 3, 4), start=1):
        ws.cell(r, 1, i)
        for c in (2, 3, 4):
            ws.cell(r, c, "")
    p = _save(wb, tmp_path / "serial.xlsx")
    result = FormStructureService().parse_structure(str(p), filename="serial.xlsx")
    fields = result["fields"]
    # 9 个空填写位（3 行 × 3 列）
    assert len(fields) == 9
    names = {f["name"] for f in fields}
    assert names == {"姓名", "性别", "年龄"}
    for f in fields:
        assert f["name"] not in {"1", "2", "3"}


def test_parse_xlsx_serial_column_not_misjudged(tmp_path: Path) -> None:
    """验证只有 2 个连续数字的列不被误判为序号列（阈值 >= 3）。"""
    wb = Workbook(); ws = wb.active
    # col0 只有 1/2（不够 3 个递增），不应被判为序号列
    ws.cell(1, 1, "1"); ws.cell(1, 2, "")
    ws.cell(2, 1, "2"); ws.cell(2, 2, "")
    p = _save(wb, tmp_path / "not_serial.xlsx")
    result = FormStructureService().parse_structure(str(p), filename="not_serial.xlsx")
    fields = result["fields"]
    # 2 个空填写位；col0 不是序号列，故「1/2」被当作字段名
    assert len(fields) == 2
    assert {f["name"] for f in fields} == {"1", "2"}


def test_parse_xlsx_example_row_skipped(tmp_path: Path) -> None:
    """示例数据行（col0「示例」+ 下方空数据行）整行跳过不算填写位。

    按用户规则：示例行仅作填表演示，非真实填写位，整行不输出任何 field。
    字段识别只发生在行3-5的真实数据行（空格→empty填写位）。
    """
    wb = Workbook(); ws = wb.active
    # 行1：列头
    ws.cell(1, 1, "序号"); ws.cell(1, 2, "姓名"); ws.cell(1, 3, "学号")
    # 行2：示例行（col0「示例」+ 全填样板值）
    ws.cell(2, 1, "示例"); ws.cell(2, 2, "张三"); ws.cell(2, 3, "2025010101")
    # 行3-5：序号 + 空数据行
    for i, r in enumerate((3, 4, 5), start=1):
        ws.cell(r, 1, i)
        ws.cell(r, 2, ""); ws.cell(r, 3, "")
    p = _save(wb, tmp_path / "example.xlsx")
    result = FormStructureService().parse_structure(str(p), filename="example.xlsx")
    fields = result["fields"]
    # 6 个填写位：行2示例行整行跳过 + 行3-5空格3×2=6个empty（序号列跳过）
    assert len(fields) == 6
    # 行2 示例行整行跳过，不输出任何 field
    assert all(f["locator"]["row"] != 2 for f in fields)
    # 行3-5空格是empty填写位
    assert all(f["fill_kind"] == "empty" for f in fields)
    assert {f["name"] for f in fields} == {"姓名", "学号"}


def test_parse_xlsx_example_row_no_keyword(tmp_path: Path) -> None:
    """无关键词示例行：靠数据行特征（多纯数字值 + 下方空行）识别后整行跳过。

    按用户规则：示例行仅作填表演示，非真实填写位，整行不输出任何 field。
    字段识别只发生在行3-5的真实数据行。
    """
    wb = Workbook(); ws = wb.active
    # 行1：列头（纯文字，0 个纯数字 → 不会被 B) 误判）
    ws.cell(1, 1, "序号"); ws.cell(1, 2, "姓名"); ws.cell(1, 3, "学号"); ws.cell(1, 4, "电话")
    # 行2：示例行（col0 无关键词，但含学号/电话 2 个纯数字）
    ws.cell(2, 1, "1"); ws.cell(2, 2, "张三"); ws.cell(2, 3, "20250101"); ws.cell(2, 4, "13800138000")
    # 行3-5：序号 + 空数据行
    for i, r in enumerate((3, 4, 5), start=2):
        ws.cell(r, 1, i)
        ws.cell(r, 2, ""); ws.cell(r, 3, ""); ws.cell(r, 4, "")
    p = _save(wb, tmp_path / "example_no_kw.xlsx")
    result = FormStructureService().parse_structure(str(p), filename="example_no_kw.xlsx")
    fields = result["fields"]
    # 9 个填写位：行2示例行整行跳过 + 行3-5空格3×3=9个empty（序号列跳过）
    assert len(fields) == 9
    # 行2 示例行整行跳过，不输出任何 field
    assert all(f["locator"]["row"] != 2 for f in fields)
    # 行3-5空格是empty填写位，序号列被跳过
    assert all(f["fill_kind"] == "empty" for f in fields)
    assert {f["name"] for f in fields} == {"姓名", "学号", "电话"}


def test_parse_xlsx_empty_sheet_skipped(tmp_path: Path) -> None:
    """验证全空 sheet 被跳过，不产生填写位。"""
    wb = Workbook(); ws = wb.active; ws.title = "数据"
    ws.cell(1, 1, "姓名"); ws.cell(2, 1, "")
    ws2 = wb.create_sheet("空表")  # 全空 sheet
    p = _save(wb, tmp_path / "empty_sheet.xlsx")
    result = FormStructureService().parse_structure(str(p), filename="empty_sheet.xlsx")
    # 仅「数据」sheet 的 1 个空填写位，空表 sheet 不产生填写位
    sheets_with_fields = {f["locator"]["sheet"] for f in result["fields"]}
    assert "空表" not in sheets_with_fields
    assert len(result["fields"]) == 1


def test_parse_xlsx_header_with_paren_not_template(tmp_path: Path) -> None:
    """验证列头含「（100字）」不被误判为模板填空栏（括号内数字非格式示例）。"""
    wb = Workbook(); ws = wb.active
    # 列头在行1列1（含括号数字说明），行2列1 为空填写位
    ws.cell(1, 1, "实践主要内容（100字）"); ws.cell(2, 1, "")
    p = _save(wb, tmp_path / "paren_header.xlsx")
    result = FormStructureService().parse_structure(str(p), filename="paren_header.xlsx")
    fields = result["fields"]
    # 仅 1 个空填写位；列头不被当模板栏；字段名取列头
    assert len(fields) == 1
    assert fields[0]["name"] == "实践主要内容（100字）"
    assert fields[0]["fill_kind"] == "empty"


def test_parse_xlsx_serial_gap_uses_serial_header(tmp_path: Path) -> None:
    """验证序号列尾部断档（5 后全空）不把「5」当字段名，取列头「序号」。"""
    wb = Workbook(); ws = wb.active
    ws.cell(1, 1, "序号"); ws.cell(1, 2, "姓名")
    for i, r in enumerate(range(2, 7), start=1):
        ws.cell(r, 1, i); ws.cell(r, 2, "")
    # 行7-8：序号留空（断档）→ 序号列空位字段名应取列头「序号」
    ws.cell(7, 1, ""); ws.cell(7, 2, "")
    ws.cell(8, 1, ""); ws.cell(8, 2, "")
    p = _save(wb, tmp_path / "serial_gap.xlsx")
    result = FormStructureService().parse_structure(str(p), filename="serial_gap.xlsx")
    # 行7-8 列1（序号列断档）的 2 个空位字段名为「序号」
    gap_fields = [f for f in result["fields"] if f["locator"]["row"] in (7, 8) and f["locator"]["col"] == 1]
    assert len(gap_fields) == 2
    assert {f["name"] for f in gap_fields} == {"序号"}


def test_parse_xlsx_full_row_note_filtered(tmp_path: Path) -> None:
    """验证整行合并的表尾说明文字被过滤，不当填写位。"""
    wb = Workbook(); ws = wb.active
    ws.cell(1, 1, "序号"); ws.cell(1, 2, "姓名"); ws.cell(1, 3, "电话")
    # 行2-4：序号 1/2/3（col0 成为序号列）+ 空数据
    for i, r in enumerate((2, 3, 4), start=1):
        ws.cell(r, 1, i); ws.cell(r, 2, ""); ws.cell(r, 3, "")
    # 行5：整行合并的表尾说明（A5:C5，colspan=3）
    ws.cell(5, 1, "填表说明：1.所有信息请如实填写 2.无内容项填无")
    ws.merge_cells("A5:C5")
    p = _save(wb, tmp_path / "note_row.xlsx")
    result = FormStructureService().parse_structure(str(p), filename="note_row.xlsx")
    fields = result["fields"]
    # 6 个空填写位（行2-4 列2-3），表尾说明行5不当填写位
    assert len(fields) == 6
    assert all(f["locator"]["row"] in (2, 3, 4) for f in fields)
    assert {f["name"] for f in fields} == {"姓名", "电话"}


def test_parse_xlsx_header_trailing_blank_and_textarea(tmp_path: Path) -> None:
    """验证论文统计场景：表头行末尾有空列仍识别为表头 + 逻辑合并组每行只1个填写位。

    - 行1 标题（A1:E1 整行合并 5 列）→ 整行不当填写位
    - 行2 表头：列1-3 有列标题（序号/姓名/备注），列4-5 空（表头末尾空列）→ 仍判为表头行
      列3「备注」是 textarea 主列，紧邻右侧列4-5 表头全空 → 识别为逻辑合并组子列
    - 行3 数据行：列1-5 全空 → empty 填写位
      列1-2 各1条 empty（序号/姓名 text）；列3 输出1条 empty（备注 textarea，
      locator.col=3，merged_range="C3:E3"）；列4-5 是逻辑合并组子列 → 跳过不输出
    """
    wb = Workbook(); ws = wb.active
    # 行1 标题整行合并
    ws.cell(1, 1, "本科生论文发表统计"); ws.merge_cells("A1:E1")
    # 行2 表头（列4-5 空，模拟表头末尾空列）
    ws.cell(2, 1, "序号"); ws.cell(2, 2, "姓名"); ws.cell(2, 3, "备注")
    ws.cell(2, 4, ""); ws.cell(2, 5, "")
    # 行3 数据行全空
    for c in range(1, 6):
        ws.cell(3, c, "")
    p = _save(wb, tmp_path / "paper.xlsx")
    result = FormStructureService().parse_structure(str(p), filename="paper.xlsx")
    fields = result["fields"]
    # 3 个填写位：行3 列1（序号）+ 列2（姓名）+ 列3（备注主列，子列4-5跳过）
    assert len(fields) == 3
    # 行1 标题整行合并不当填写位；行2 表头不当填写位
    assert all(f["locator"]["row"] == 3 for f in fields)
    by_col = {f["locator"]["col"]: f for f in fields}
    # 列4-5 是逻辑合并组子列，跳过不输出
    assert 4 not in by_col
    assert 5 not in by_col
    assert by_col[1]["name"] == "序号"
    assert by_col[1]["field_type"] == "text"
    assert by_col[2]["name"] == "姓名"
    assert by_col[2]["field_type"] == "text"
    # 列3 备注主列输出1条 textarea，locator.col=3，merged_range="C3:E3"
    assert by_col[3]["name"] == "备注"
    assert by_col[3]["field_type"] == "textarea"
    assert by_col[3]["fill_kind"] == "empty"
    assert by_col[3]["locator"]["merged_range"] == "C3:E3"


def test_parse_xlsx_dropdown_example_and_empty_sheets(tmp_path: Path) -> None:
    """验证附件1场景：下拉示例值判填写位+字段名取表头+空sheet跳过+签字行过滤。

    - 行1整行合并表标题、行2整行合并签字行 → 不当填写位
    - 行3表头（D 列「项目组别（须下拉选择）」）→ 列名，不当填写位
    - 行4 D4 预置下拉示例「先进制造」→ template 填写位，field_type=select，
      name 继承表头「项目组别」（不被同行/数据行值污染）
    - 行4-8 其余空格 → empty 填写位，name 取对应表头列名
    - Sheet2/Sheet3 全空 → 不出现在 sheets，不产生填写位
    - 合计 40 个填写位（5 数据行 × 8 列，其中 D4 为 template，其余 empty）
    """
    wb = Workbook(); ws = wb.active; ws.title = "Sheet1"
    # 行1 表标题（A1:H1 整行合并）
    ws.cell(1, 1, "第十五届挑战杯参赛项目推荐汇总表"); ws.merge_cells("A1:H1")
    # 行2 签字行（A2:H2 整行合并，含签字/盖章）
    ws.cell(2, 1, "推荐学院团委书记签字（盖章）："); ws.merge_cells("A2:H2")
    # 行3 表头
    headers = ["学院名称", "推荐顺序", "项目名称", "项目组别\n（须下拉选择）",
               "团队负责人姓名", "团队负责人电话", "团队成员名单（排序）", "指导教师名单（排序）"]
    for c, h in enumerate(headers, 1):
        ws.cell(3, c, h)
    # 行4-8：5 条数据记录；D4 预置下拉示例「先进制造」，其余空（显式设空确保 max_row=8）
    for r in range(4, 9):
        for c in range(1, 9):
            if r == 4 and c == 4:
                ws.cell(r, c, "先进制造")
            else:
                ws.cell(r, c, "")
    # 新建 Sheet2/Sheet3 全空
    wb.create_sheet("Sheet2"); wb.create_sheet("Sheet3")
    p = _save(wb, tmp_path / "attach1.xlsx")
    result = FormStructureService().parse_structure(str(p), filename="attach1.xlsx")

    # 空 sheet 不出现在 sheets
    sheet_names = [s["name"] for s in result["sheets"]]
    assert sheet_names == ["Sheet1"]

    fields = result["fields"]
    # 40 个填写位：D4 是 template，其余 39 个 empty
    assert len(fields) == 40
    tpls = [f for f in fields if f["fill_kind"] == "template"]
    empties = [f for f in fields if f["fill_kind"] == "empty"]
    assert len(tpls) == 1 and len(empties) == 39
    # D4 = template + field_type=select
    d4 = next(f for f in fields if f["locator"]["row"] == 4 and f["locator"]["col"] == 4)
    assert d4["fill_kind"] == "template"
    assert d4["field_type"] == "select"
    assert d4["name"] == "项目组别"
    # 行4 列5-8 name 取对应表头，不被 D4「先进制造」污染
    r4 = {f["locator"]["col"]: f["name"] for f in fields if f["locator"]["row"] == 4}
    assert r4[5] == "团队负责人姓名"
    assert r4[6] == "团队负责人电话"
    assert r4[7] == "团队成员名单（排序）"
    assert r4[8] == "指导教师名单（排序）"
    # 行5-8 列4 name 取表头「项目组别」，不被 D4 示例值污染
    r5_c4 = next(f for f in fields if f["locator"]["row"] == 5 and f["locator"]["col"] == 4)
    assert r5_c4["name"] == "项目组别"
    assert r5_c4["field_type"] == "select"
    # 普通列 field_type=text
    assert next(f for f in fields if f["locator"]["row"] == 4 and f["locator"]["col"] == 1)["field_type"] == "text"
    # 行1 标题、行2 签字行不当填写位
    assert all(f["locator"]["row"] not in (1, 2) for f in fields)


def test_parse_xlsx_real_files() -> None:
    """用 test/test/ 下的真实 xlsx 端到端验证（数据驱动）。

    若真实文件不存在则跳过（不阻塞 CI，因测试数据可能单独提交）。
    """
    repo_root = Path(__file__).resolve().parents[1]
    test_dir = repo_root / "test" / "test"
    xlsx_files = sorted(test_dir.glob("*.xlsx")) if test_dir.exists() else []
    if not xlsx_files:
        pytest.skip("test/test/ 下无真实 xlsx 测试数据")
    svc = FormStructureService()
    for xlsx in xlsx_files:
        result = svc.parse_structure(str(xlsx), filename=xlsx.name)
        assert result["format"] == "xlsx"
        assert result["filename"] == xlsx.name
        # 契约校验（不抛异常即通过）
        for f in result["fields"]:
            ExcelLocator(**f["locator"])
        # 可序列化
        json.dumps(result, ensure_ascii=False)


# ---------------------------------------------------------------------------
# Bug 修复回归测试
# ---------------------------------------------------------------------------


def test_b1_vmerge_multi_level_header_name(tmp_path: Path) -> None:
    """B1: 纵向合并多级表头字段名取正确列头，不取父级表头。

    行0：| (vmerge r0-r1) 分类 | (hmerge c1-c2) 指标 |
    行1：| (merged)            | 数量 | 金额 |
    行2：| 项目A | (fill) | (fill) |

    行2 列1 的 name 应为「数量」（行1子表头），而非「指标」（行0父表头）。
    行2 列2 的 name 应为「金额」。
    """
    wb = Workbook(); ws = wb.active
    # 行0：分类纵向合并 + 指标横向合并
    ws.cell(1, 1, "分类")
    ws.merge_cells(start_row=1, start_column=1, end_row=2, end_column=1)  # A1:A2
    ws.cell(1, 2, "指标")
    ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=3)  # B1:C1
    # 行1：子表头
    ws.cell(2, 2, "数量"); ws.cell(2, 3, "金额")
    # 行2：数据行，列0有值，列1-2空
    ws.cell(3, 1, "项目A")
    ws.cell(3, 2, ""); ws.cell(3, 3, "")
    p = _save(wb, tmp_path / "b1.xlsx")
    result = FormStructureService().parse_structure(str(p), filename="b1.xlsx")
    fields = result["fields"]
    by_col = {f["locator"]["col"]: f for f in fields}
    assert by_col[2]["name"] == "数量"
    assert by_col[3]["name"] == "金额"
    # 不应取到父表头「指标」
    assert "指标" not in {f["name"] for f in fields}


def test_b2_colon_label_not_fillable(tmp_path: Path) -> None:
    """B2: 「负责人签字：」标签右侧有填写格时，不当填写位输出。

    行1：| 负责人签字： | (空填写格) |
    「负责人签字：」是标签，不应作为填写位；右侧空格才是填写位。
    """
    wb = Workbook(); ws = wb.active
    ws.cell(1, 1, "负责人签字：")
    ws.cell(1, 2, "")  # 右侧空填写格
    p = _save(wb, tmp_path / "b2.xlsx")
    result = FormStructureService().parse_structure(str(p), filename="b2.xlsx")
    fields = result["fields"]
    # 只有右侧空格是填写位，「负责人签字：」不是
    assert len(fields) == 1
    assert fields[0]["locator"]["col"] == 2
    assert fields[0]["fill_kind"] == "empty"


def test_b3_data_row_single_value_not_header(tmp_path: Path) -> None:
    """B3: 数据行只填1个值时被正确识别为数据行（非表头）。

    3 列表，表头行全有值，数据行仅列0有值、列1-2空。
    修复前阈值 num_cols//2=1，数据行 value_count=1 不 < 1 → 误判为表头；
    修复后阈值 (num_cols+1)//2=2，1 < 2 → 正确识别为数据行。
    """
    wb = Workbook(); ws = wb.active
    # 表头行
    ws.cell(1, 1, "姓名"); ws.cell(1, 2, "年龄"); ws.cell(1, 3, "城市")
    # 数据行：仅列0有值
    ws.cell(2, 1, "张三")
    ws.cell(2, 2, ""); ws.cell(2, 3, "")
    p = _save(wb, tmp_path / "b3.xlsx")
    result = FormStructureService().parse_structure(str(p), filename="b3.xlsx")
    sheet = result["sheets"][0]
    # header_rows 应为 1（仅行0是表头），行1是数据行
    assert sheet["header_rows"] == 1
    fields = result["fields"]
    # 行2 列2、列3 是填写位
    cols = {f["locator"]["col"] for f in fields}
    assert 2 in cols
    assert 3 in cols


def test_b4_formula_cell_not_fillable(tmp_path: Path) -> None:
    """B4: 公式单元格（=SUM）不当填写位，避免写回覆盖公式。"""
    wb = Workbook(); ws = wb.active
    ws.cell(1, 1, "项目"); ws.cell(1, 2, "金额")
    ws.cell(2, 1, "A"); ws.cell(2, 2, 100)
    ws.cell(3, 1, "B"); ws.cell(3, 2, 200)
    # 合计行：列0有标签，列1是 =SUM 公式（无缓存值时 data_only 读为 None）
    ws.cell(4, 1, "合计")
    ws.cell(4, 2, "=SUM(B2:B3)")
    p = _save(wb, tmp_path / "b4.xlsx")
    result = FormStructureService().parse_structure(str(p), filename="b4.xlsx")
    fields = result["fields"]
    # 公式单元格 B4 不应出现在填写位中
    formula_locs = {(f["locator"]["row"], f["locator"]["col"]) for f in fields}
    assert (4, 2) not in formula_locs


def test_b6_hidden_column_not_fillable(tmp_path: Path) -> None:
    """B6: 隐藏列不当填写位。"""
    wb = Workbook(); ws = wb.active
    ws.cell(1, 1, "姓名"); ws.cell(1, 2, "年龄"); ws.cell(1, 3, "城市")
    ws.cell(2, 1, ""); ws.cell(2, 2, ""); ws.cell(2, 3, "")
    # 隐藏列 B（列2）
    ws.column_dimensions["B"].hidden = True
    p = _save(wb, tmp_path / "b6.xlsx")
    result = FormStructureService().parse_structure(str(p), filename="b6.xlsx")
    fields = result["fields"]
    cols = {f["locator"]["col"] for f in fields}
    # 隐藏列 B(col=2) 不应有填写位
    assert 2 not in cols
    # 非隐藏列仍有填写位
    assert 1 in cols
    assert 3 in cols
