from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.main import app
from app.modules.task_semantics.service import TaskSemanticsService


@pytest.fixture()
def client(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> TestClient:
    monkeypatch.setenv("AFF_DATA_DIR", str(tmp_path))
    return TestClient(app)


def test_cell_value_strips_explanation() -> None:
    from app.api.forms import _cell_value

    assert (
        _cell_value("根据您提供的材料，第2行基本档案摘要中记录的本人姓名是**张三** [1]", True)
        == "张三"
    )
    assert (
        _cell_value("根据材料，张三（本人）的移动电话（手机号）为 13800138000 [1]", True)
        == "13800138000"
    )
    assert (
        _cell_value("2016-09 ~ 2020-06### References- 学业与在校实践.xlsx", True)
        == "2016-09 ~ 2020-06"
    )
    assert _cell_value("无", True) is None
    assert _cell_value("张三", False) is None


def test_aligns_talent_form_labels() -> None:
    structure = {
        "sheets": [
            {
                "name": "信息项",
                "cells": [
                    {"row": 3, "col": 1, "text": "手机号码"},
                    {"row": 3, "col": 2, "text": ""},
                    {"row": 12, "col": 1, "text": "最近岗位"},
                    {"row": 14, "col": 1, "text": "专业证书（可多项）"},
                    {"row": 16, "col": 1, "text": "血型"},
                ],
            }
        ],
        "fields": [
            {
                "id": "phone",
                "name": "填写内容",
                "locator": {"kind": "excel_cell", "sheet": "信息项", "row": 3, "col": 2},
            },
            {
                "id": "role",
                "name": "填写内容",
                "locator": {"kind": "excel_cell", "sheet": "信息项", "row": 12, "col": 2},
            },
            {
                "id": "cert",
                "name": "填写内容",
                "locator": {"kind": "excel_cell", "sheet": "信息项", "row": 14, "col": 2},
            },
            {
                "id": "blood",
                "name": "填写内容",
                "locator": {"kind": "excel_cell", "sheet": "信息项", "row": 16, "col": 2},
            },
            {
                "id": "school",
                "name": "学校/院系",
                "locator": {"kind": "excel_cell", "sheet": "在校经历", "row": 2, "col": 2},
            },
        ],
    }
    spec = TaskSemanticsService().build_task_spec(structure)
    by_id = {item["field_id"]: item for item in spec["fields"]}
    assert by_id["phone"]["status"] == "aligned"
    assert by_id["phone"]["attribute"] == "手机号"
    assert by_id["phone"]["label"] == "手机号码"
    assert by_id["role"]["entity"] == "工作经历"
    assert by_id["role"]["attribute"] == "职务"
    assert by_id["cert"]["attribute"] == "证书名称"
    assert by_id["blood"]["status"] == "unaligned"
    assert by_id["blood"]["query"] is None
    assert by_id["school"]["attribute"] == "学校"
    assert "第3行" in by_id["phone"]["query"]


def test_fill_writes_only_evidenced_values(
    client, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from aff_contracts import RagContext, RagQueryResult

    class FakeRag:
        async def aquery_value(self, query: str, *, attribute: str | None = None) -> RagQueryResult:
            if "手机号" in query:
                return RagQueryResult(
                    answer="13800138000",
                    contexts=[RagContext(content="手机号码 13800138000", doc_id="简介.txt", score=0.9)],
                )
            return RagQueryResult(answer="无", contexts=[])

    monkeypatch.setattr("app.api.forms.get_rag_service", lambda: FakeRag())
    book = Workbook()
    sheet = book.active
    sheet.title = "信息项"
    sheet["A1"] = "采集项"
    sheet["B1"] = "填写内容"
    sheet["A2"] = "手机号码"
    sheet["A3"] = "血型"
    path = tmp_path / "表.xlsx"
    book.save(path)
    uploaded = client.post(
        "/api/forms/upload",
        files={
            "file": (
                "表.xlsx",
                path.read_bytes(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert uploaded.status_code == 201, uploaded.text
    job_id = uploaded.json()["id"]
    uploaded_body = uploaded.json()
    phone = next(field for field in uploaded_body["fields"] if field["name"] == "手机号码")
    blood = next(field for field in uploaded_body["fields"] if field["name"] == "血型")
    phone_res = client.post(f"/api/forms/jobs/{job_id}/fill?field_id={phone['id']}")
    assert phone_res.status_code == 200, phone_res.text
    assert phone_res.json()["field"]["value"] == "13800138000"
    assert phone_res.json()["field"]["status"] == "suggested"
    blood_res = client.post(f"/api/forms/jobs/{job_id}/fill?field_id={blood['id']}")
    assert blood_res.status_code == 200, blood_res.text
    assert blood_res.json()["field"]["value"] is None
    assert blood_res.json()["field"]["notes"] == "未对齐"
    saved = client.get(f"/api/forms/jobs/{job_id}").json()
    by_name = {field["name"]: field for field in saved["fields"]}
    assert by_name["手机号码"]["value"] == "13800138000"
    assert by_name["血型"]["value"] is None


def test_patch_can_add_and_delete_fields(client, tmp_path: Path) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "信息项"
    sheet["A1"] = "采集项"
    sheet["B1"] = "填写内容"
    sheet["A2"] = "手机号码"
    path = tmp_path / "表.xlsx"
    book.save(path)
    uploaded = client.post(
        "/api/forms/upload",
        files={
            "file": (
                "表.xlsx",
                path.read_bytes(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert uploaded.status_code == 201, uploaded.text
    job_id = uploaded.json()["id"]
    added = client.patch(
        f"/api/forms/jobs/{job_id}/fields",
        json={
            "fields": [
                {
                    "id": "manual-email",
                    "name": "邮箱",
                    "value": None,
                    "status": "empty",
                    "locator": {"kind": "word_bookmark", "name": "manual"},
                }
            ]
        },
    )
    assert added.status_code == 200, added.text
    body = added.json()
    assert [field["id"] for field in body["fields"]] == ["manual-email"]
    assert body["fields"][0]["notes"] == "本人.邮箱"
    assert body["task_spec"]["fields"][0]["status"] == "aligned"
    deleted = client.patch(f"/api/forms/jobs/{job_id}/fields", json={"fields": []})
    assert deleted.status_code == 200
    assert deleted.json()["fields"] == []


def test_export_writes_suggested_value_back_to_xlsx(client, tmp_path: Path) -> None:
    from openpyxl import load_workbook

    book = Workbook()
    sheet = book.active
    sheet.title = "信息项"
    sheet["A1"] = "采集项"
    sheet["B1"] = "填写内容"
    sheet["A2"] = "手机号码"
    sheet["A3"] = "血型"
    path = tmp_path / "表.xlsx"
    book.save(path)
    uploaded = client.post(
        "/api/forms/upload",
        files={
            "file": (
                "表.xlsx",
                path.read_bytes(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert uploaded.status_code == 201, uploaded.text
    job = uploaded.json()
    fields = job["fields"]
    for field in fields:
        if field["name"] == "手机号码":
            field["value"] = "13800138000"
            field["status"] = "suggested"
    patched = client.patch(f"/api/forms/jobs/{job['id']}/fields", json={"fields": fields})
    assert patched.status_code == 200, patched.text
    exported = client.post(f"/api/forms/jobs/{job['id']}/export")
    assert exported.status_code == 200, exported.text
    assert exported.json()["written"] == 1
    downloaded = client.get(f"/api/forms/jobs/{job['id']}/download")
    assert downloaded.status_code == 200, downloaded.text
    assert ".filled.xlsx" in downloaded.headers["content-disposition"]
    out = tmp_path / "out.xlsx"
    out.write_bytes(downloaded.content)
    saved = load_workbook(out)
    assert saved["信息项"]["B2"].value == "13800138000"
    assert saved["信息项"]["B3"].value is None


def test_upload_xlsx_stores_task_spec(client, tmp_path) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "信息项"
    sheet["A1"] = "采集项"
    sheet["B1"] = "填写内容"
    sheet["A2"] = "姓名（与身份证一致）"
    sheet["B2"] = None
    path = tmp_path / "人才信息采集表.xlsx"
    book.save(path)
    files = {
        "file": (
            "人才信息采集表.xlsx",
            path.read_bytes(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    }
    res = client.post("/api/forms/upload", files=files)
    assert res.status_code == 201, res.text
    body = res.json()
    names = [field["name"] for field in body["fields"]]
    assert "姓名（与身份证一致）" in names
    matched = next(field for field in body["fields"] if field["name"] == "姓名（与身份证一致）")
    assert matched["notes"] == "本人.姓名"
    assert matched["value"] is None
    assert body["task_spec"]["schema_ref"] == "personal_kb@0.1.0"
