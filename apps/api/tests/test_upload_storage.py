from __future__ import annotations

import json
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

from app.main import app
from app.modules.orchestrator import storage


@pytest.fixture()
def client(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> TestClient:
    monkeypatch.setenv("AFF_DATA_DIR", str(tmp_path))
    # reset cached settings if any — AppConfig reads env each get_config()
    return TestClient(app)


def test_knowledge_upload_list_delete(client: TestClient, tmp_path: Path) -> None:
    files = {"file": ("简历.pdf", b"%PDF-1.4 mock content", "application/pdf")}
    res = client.post("/api/knowledge/upload", files=files)
    assert res.status_code == 201, res.text
    body = res.json()
    assert body["filename"] == "简历.pdf"
    assert body["size"] == len(b"%PDF-1.4 mock content")
    assert body["status"] == "stored"
    assert body["sha256"]
    doc_id = body["id"]

    # file on disk under data/knowledge/raw/{id}/
    raw_dir = tmp_path / "knowledge" / "raw" / doc_id
    assert raw_dir.is_dir()
    assert (raw_dir / "简历.pdf").read_bytes().startswith(b"%PDF")

    listed = client.get("/api/knowledge")
    assert listed.status_code == 200
    assert any(x["id"] == doc_id for x in listed.json())

    got = client.get(f"/api/knowledge/{doc_id}")
    assert got.status_code == 200
    assert got.json()["path"].endswith(f"knowledge/raw/{doc_id}/简历.pdf") or "简历.pdf" in got.json()["path"]

    deleted = client.delete(f"/api/knowledge/{doc_id}")
    assert deleted.status_code == 200
    assert not raw_dir.exists()
    assert client.get(f"/api/knowledge/{doc_id}").status_code == 404


def test_knowledge_upload_ingests_and_writes_graph(
    client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    from aff_contracts import UpsertDocumentResult

    class FakeRag:
        def __init__(self) -> None:
            self.modes: list[str] = []

        async def aupsert(self, request):
            self.modes.append(request.mode)
            assert request.bundle.text == "姓名：张三"
            return UpsertDocumentResult(
                doc_id=request.bundle.doc_id,
                status="ready",
                stats={"mode": request.mode},
            )

    fake = FakeRag()
    monkeypatch.setattr(
        "app.api.knowledge.get_rag_service",
        lambda: fake,
    )
    files = {"file": ("简介.txt", "姓名：张三".encode("utf-8"), "text/plain")}
    res = client.post("/api/knowledge/upload", files=files)
    assert res.status_code == 201, res.text
    body = res.json()
    assert body["ingest_status"] == "ready"
    assert body["graph_status"] == "ready"
    assert body["bundle"]["text"] == "姓名：张三"
    assert body["bundle"]["doc_id"] == body["id"]
    assert body["note"] == "已写入图谱"
    assert fake.modes == ["insert"]


def test_knowledge_query_uses_graph(
    client: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    from aff_contracts import RagContext, RagQueryResult

    class FakeRag:
        async def aquery(self, request):
            assert request.query == "张三的手机号"
            assert request.mode == "mix"
            return RagQueryResult(
                answer="13800138000",
                contexts=[RagContext(content="手机号码 13800138000", doc_id="简介.txt", score=0.8)],
            )

    monkeypatch.setattr("app.api.knowledge.get_rag_service", lambda: FakeRag())
    res = client.post("/api/knowledge/query", json={"query": "张三的手机号"})
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["answer"] == "13800138000"
    assert body["contexts"][0]["doc_id"] == "简介.txt"


def test_forms_upload_keeps_original_bytes(client: TestClient, tmp_path: Path) -> None:
    payload = b"PK\x03\x04 fake-docx-bytes"
    files = {
        "file": (
            "入职表.docx",
            payload,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }
    res = client.post("/api/forms/upload", files=files)
    assert res.status_code == 201, res.text
    body = res.json()
    job_id = body["id"]
    assert body["format"] == "docx"
    assert body["step"] == "uploaded"
    assert body["fields"] == []

    saved = tmp_path / "forms" / "raw" / job_id / "入职表.docx"
    assert saved.read_bytes() == payload

    jobs = client.get("/api/forms/jobs")
    assert any(j["id"] == job_id for j in jobs.json())

    deleted = client.delete(f"/api/forms/jobs/{job_id}")
    assert deleted.status_code == 200
    assert not saved.exists()
    assert client.get(f"/api/forms/jobs/{job_id}").status_code == 404


def test_forms_upload_parses_docx_fields(client: TestClient) -> None:
    from io import BytesIO

    from docx import Document

    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "姓名"
    table.cell(0, 1).text = ""
    buf = BytesIO()
    doc.save(buf)
    files = {
        "file": (
            "登记表.docx",
            buf.getvalue(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    }
    res = client.post("/api/forms/upload", files=files)
    assert res.status_code == 201, res.text
    body = res.json()
    assert body["step"] == "parsed"
    names = [field["name"] for field in body["fields"]]
    assert "姓名" in names
    assert body["fields"][0]["status"] == "empty"
    assert body["fields"][0]["value"] is None
    assert body["structure_notes"]


def test_reject_empty_file(client: TestClient) -> None:
    files = {"file": ("empty.txt", b"", "text/plain")}
    res = client.post("/api/knowledge/upload", files=files)
    assert res.status_code == 400


def test_index_json_written(client: TestClient, tmp_path: Path) -> None:
    files = {"file": ("a.txt", b"hello", "text/plain")}
    client.post("/api/knowledge/upload", files=files)
    index = tmp_path / "knowledge" / "index.json"
    assert index.is_file()
    data = json.loads(index.read_text(encoding="utf-8"))
    assert isinstance(data, list) and data[0]["filename"] == "a.txt"
