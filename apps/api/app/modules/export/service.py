from __future__ import annotations

import shutil
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell

from app.core.paths import exports_dir


_WRITABLE = {"suggested", "confirmed", "manual"}


def filled_name(filename: str) -> str:
    path = Path(filename or "form.bin")
    return f"{path.stem}.filled{path.suffix}"


def write_filled_copy(source: Path, filename: str, fields: list[dict[str, Any]], dest_dir: Path | None = None) -> tuple[Path, int]:
    suffix = source.suffix.lower()
    if suffix not in {".docx", ".xlsx"}:
        raise ValueError("当前只回填 .docx 和 .xlsx")
    target_dir = dest_dir or exports_dir()
    target_dir.mkdir(parents=True, exist_ok=True)
    dest = target_dir / filled_name(filename or source.name)
    shutil.copy2(source, dest)
    written = 0
    if suffix == ".xlsx":
        written = _write_xlsx(dest, fields)
    else:
        written = _write_docx(dest, fields)
    return dest, written


def _should_write(field: dict[str, Any]) -> bool:
    value = field.get("value")
    if value is None or str(value).strip() == "":
        return False
    return str(field.get("status") or "") in _WRITABLE


def _write_xlsx(path: Path, fields: list[dict[str, Any]]) -> int:
    book = load_workbook(path)
    written = 0
    for field in fields:
        if not _should_write(field):
            continue
        locator = field.get("locator") if isinstance(field.get("locator"), dict) else {}
        if locator.get("kind") != "excel_cell":
            continue
        sheet = str(locator.get("sheet") or "")
        row = locator.get("row")
        col = locator.get("col")
        if sheet not in book.sheetnames or not isinstance(row, int) or not isinstance(col, int):
            continue
        _set_excel_cell(book[sheet], row, col, str(field.get("value")))
        written += 1
    book.save(path)
    return written


def _set_excel_cell(sheet: Any, row: int, col: int, value: str) -> None:
    cell = sheet.cell(row, col)
    if isinstance(cell, MergedCell):
        for merged in sheet.merged_cells.ranges:
            if cell.coordinate in merged:
                sheet.cell(merged.min_row, merged.min_col).value = value
                return
        return
    cell.value = value


def _write_docx(path: Path, fields: list[dict[str, Any]]) -> int:
    from docx import Document

    document = Document(path)
    written = 0
    for field in fields:
        if not _should_write(field):
            continue
        locator = field.get("locator") if isinstance(field.get("locator"), dict) else {}
        if locator.get("kind") != "word_cell":
            continue
        table_index = locator.get("table_index")
        row = locator.get("row")
        col = locator.get("col")
        if not isinstance(table_index, int) or not isinstance(row, int) or not isinstance(col, int):
            continue
        if table_index < 0 or table_index >= len(document.tables):
            continue
        table = document.tables[table_index]
        if row < 0 or row >= len(table.rows) or col < 0 or col >= len(table.rows[row].cells):
            continue
        table.cell(row, col).text = str(field.get("value"))
        written += 1
    document.save(path)
    return written
